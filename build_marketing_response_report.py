from __future__ import annotations

from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
from matplotlib import font_manager
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\wangjunyi\Documents\New project")
OUTPUT_DIR = ROOT / "outputs" / "marketing_response_model"
DOCX_PATH = OUTPUT_DIR / "营销活动响应预测模型报告.docx"
CHART_PATH = OUTPUT_DIR / "feature_coefficients_chart.png"

BLUE = "1F4E79"
LIGHT_BLUE = "D9EAF7"
GREEN = "E2F0D9"
RED = "FCE4D6"
GRAY = "F2F2F2"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, color: str | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(10)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False

    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_text(cell, header, bold=True, color="FFFFFF")
        set_cell_shading(cell, BLUE)
        cell.width = Cm(widths[i])

    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
            cells[i].width = Cm(widths[i])

    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.1


def style_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.2
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size, color in [
        ("Title", 24, BLUE),
        ("Heading 1", 16, BLUE),
        ("Heading 2", 13, BLUE),
    ]:
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True


def add_callout(document: Document, title: str, body: str, fill: str = LIGHT_BLUE) -> None:
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    title_run = p.add_run(title)
    title_run.bold = True
    title_run.font.color.rgb = RGBColor.from_string(BLUE)
    title_run.font.size = Pt(11)
    p.add_run("\n" + body)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.line_spacing = 1.2


def build_chart(coefficients: pd.DataFrame) -> None:
    font_manager.fontManager.addfont(r"C:\Windows\Fonts\msyh.ttc")
    plt.rcParams["font.sans-serif"] = ["Microsoft YaHei"]
    plt.rcParams["axes.unicode_minus"] = False

    selected = pd.concat(
        [
            coefficients.sort_values("coefficient").head(6),
            coefficients.sort_values("coefficient").tail(6),
        ]
    )
    selected = selected.assign(
        label=lambda d: d["feature"]
        .str.replace("num__", "", regex=False)
        .str.replace("cat__", "", regex=False)
    ).sort_values("coefficient")

    colors = ["#C0504D" if value < 0 else "#4F81BD" for value in selected["coefficient"]]
    plt.figure(figsize=(8.2, 4.8))
    plt.barh(selected["label"], selected["coefficient"], color=colors)
    plt.axvline(0, color="#666666", linewidth=0.8)
    plt.title("关键变量对营销响应概率的影响方向", fontsize=13)
    plt.xlabel("逻辑回归标准化系数")
    plt.tight_layout()
    plt.savefig(CHART_PATH, dpi=220)
    plt.close()


def main() -> None:
    metrics = pd.read_csv(OUTPUT_DIR / "model_metrics.csv").iloc[0]
    coefficients = pd.read_csv(OUTPUT_DIR / "feature_coefficients.csv")
    scores = pd.read_csv(OUTPUT_DIR / "customer_scores.csv")
    source = pd.read_csv(r"C:\Users\wangjunyi\Desktop\marketing_campaign.csv", sep="\t")

    build_chart(coefficients)

    recommended = int((scores["RecommendContact"] == "Yes").sum())
    total = len(scores)
    response_rate = source["Response"].mean()
    threshold = 3 / 11

    document = Document()
    style_document(document)

    title = document.add_paragraph()
    title.style = document.styles["Title"]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("营销活动响应预测模型报告")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("基于客户画像、消费行为与历史活动响应的投放决策模型")
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string("666666")

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("数据集：marketing_campaign.csv    |    输出日期：2026-05-08")

    document.add_paragraph()
    add_callout(
        document,
        "核心结论",
        f"模型在测试集上的 AUC 为 {metrics['auc']:.3f}，说明对客户响应倾向具有较好的区分能力。"
        f"按单次联系成本 3、响应收益 11 计算，建议投放阈值为 {threshold:.1%}；"
        f"当前模型建议优先联系 {recommended:,} 名客户，占总客户数 {recommended / total:.1%}。",
        GREEN,
    )

    document.add_heading("1. 项目背景与目标", level=1)
    document.add_paragraph(
        "本报告基于营销活动客户数据，建立客户响应概率预测模型。模型的目标不是简单判断客户是否会响应，"
        "而是把预测概率转化为可执行的投放建议，帮助在预算有限的情况下优先触达更可能带来正向收益的客户。"
    )

    add_table(
        document,
        ["项目", "说明"],
        [
            ["样本量", f"{total:,} 条客户记录"],
            ["目标变量", "Response，1 表示客户响应营销活动，0 表示未响应"],
            ["实际响应率", f"{response_rate:.1%}"],
            ["缺失处理", "Income 有 24 条缺失，使用训练集收入中位数填补"],
            ["投放决策规则", "当 11 × 响应概率 - 3 > 0 时，建议投放"],
        ],
        [4.0, 11.0],
    )

    document.add_heading("2. 数学模型设计", level=1)
    document.add_paragraph(
        "本次采用逻辑回归作为第一版模型。该模型适合处理二分类问题，同时保留较好的可解释性，"
        "便于说明哪些客户特征会提高或降低响应概率。"
    )
    add_callout(
        document,
        "模型形式",
        "P(Response = 1) = 1 / (1 + exp(-(b0 + b1*x1 + b2*x2 + ... + bn*xn)))",
        LIGHT_BLUE,
    )
    document.add_paragraph(
        "模型输入包括客户收入、年龄、最近一次购买间隔、各品类消费金额、线上/线下/目录购买次数、"
        "历史活动接受情况、教育程度和婚姻状态等变量。数值变量经过缺失填补与标准化处理，类别变量经过独热编码处理。"
    )

    document.add_heading("3. 模型评估结果", level=1)
    add_table(
        document,
        ["指标", "结果", "解释"],
        [
            ["AUC", f"{metrics['auc']:.3f}", "越接近 1，模型越能区分响应客户与非响应客户"],
            ["准确率（0.50 阈值）", f"{metrics['accuracy_at_0_50']:.1%}", "按 50% 阈值判断时的整体正确率"],
            ["精确率（0.50 阈值）", f"{metrics['precision_at_0_50']:.1%}", "预测会响应的客户中实际响应的比例"],
            ["召回率（0.50 阈值）", f"{metrics['recall_at_0_50']:.1%}", "实际响应客户中被模型识别出来的比例"],
            ["收益阈值召回率", f"{metrics['recall_at_roi_threshold']:.1%}", "按收益阈值投放时覆盖实际响应客户的能力"],
        ],
        [4.0, 3.2, 7.8],
    )

    document.add_paragraph(
        "从业务角度看，AUC 达到 0.900，说明模型排序能力较强。由于营销响应客户本身占比只有约 14.9%，"
        "使用收益阈值比固定 50% 阈值更适合投放决策。"
    )

    document.add_heading("4. 关键影响因素", level=1)
    document.add_picture(str(CHART_PATH), width=Inches(6.3))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(
        "上图展示了逻辑回归系数较大的变量。正系数表示该变量提高响应概率，负系数表示该变量降低响应概率。"
    )
    document.add_page_break()

    positive = coefficients.sort_values("coefficient", ascending=False).head(6)
    negative = coefficients.sort_values("coefficient").head(6)
    rows = []
    for _, row in positive.iterrows():
        rows.append([
            row["feature"].replace("num__", "").replace("cat__", ""),
            f"{row['coefficient']:.3f}",
            "提高响应概率",
        ])
    for _, row in negative.iterrows():
        rows.append([
            row["feature"].replace("num__", "").replace("cat__", ""),
            f"{row['coefficient']:.3f}",
            "降低响应概率",
        ])
    add_table(document, ["变量", "系数", "方向"], rows, [6.8, 2.5, 4.8])

    document.add_heading("5. 投放策略建议", level=1)
    add_table(
        document,
        ["策略", "执行方式", "预期作用"],
        [
            [
                "优先名单投放",
                "筛选 RecommendContact = Yes 的客户，并按 ResponseProbability 从高到低排序",
                "优先覆盖高概率客户，提升预算使用效率",
            ],
            [
                "预算受限场景",
                "若预算不足以覆盖全部 946 名建议客户，可截取概率排名前 20%-30%",
                "控制成本，同时保留主要响应机会",
            ],
            [
                "客户分层运营",
                "高概率客户直接促销；中概率客户使用低成本触达；低概率客户暂缓投放",
                "减少无效触达，降低营销疲劳",
            ],
            [
                "持续迭代",
                "下一轮活动结束后，把真实响应结果追加到数据集中重新训练",
                "让模型随客户行为变化而更新",
            ],
        ],
        [3.2, 7.0, 4.5],
    )

    section = document.add_section(WD_SECTION.NEW_PAGE)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    document.add_heading("附录：输出文件说明", level=1)
    add_table(
        document,
        ["文件", "用途"],
        [
            ["customer_scores.csv", "客户级预测结果，包含响应概率、预期利润和是否建议投放"],
            ["feature_coefficients.csv", "模型变量系数，用于解释变量影响方向"],
            ["model_metrics.csv", "模型评估指标"],
            ["模型文件（.joblib）", "可复用的训练后模型文件"],
            ["marketing_response_model.py", "模型训练与预测脚本"],
        ],
        [5.8, 8.7],
    )

    document.add_heading("附录：方法限制", level=1)
    document.add_paragraph(
        "本模型基于历史营销活动数据建立，适合辅助投放决策，但不应视为绝对判断。"
        "如果未来活动形式、优惠力度、客户群体或市场环境发生变化，模型效果可能下降。"
        "建议在每次活动后使用最新反馈数据重新训练，并与小规模 A/B 测试结合验证实际收益。"
    )

    for section in document.sections:
        section.footer.is_linked_to_previous = False
        footer = section.footer.paragraphs[0]
        footer.clear()
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("营销活动响应预测模型报告")

    document.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
